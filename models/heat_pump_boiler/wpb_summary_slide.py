"""Single-slide summary: method + the classified-vs-GWR comparison.

    python -m models.heat_pump_boiler.wpb_summary_slide

Standalone from wpb_slides.py (which builds the full deck); this produces
one slide for use on its own, e.g. dropped into another deck.
"""
from __future__ import annotations

import json

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from . import config as C
from .wpb_comparison_chart import build as build_chart
from .wpb_comparison_chart import OUT as CHART_PNG

DECK = C.ARTIFACTS / "wpb_summary_slide.pptx"
INK, ACCENT, MUTED, WARN = 0x1F2937, 0x0F766E, 0x6B7280, 0xB45309


def _rgb(v):
    return RGBColor(v >> 16 & 255, v >> 8 & 255, v & 255)


def _header(shapes, text, x, y, w):
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.32))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size, p.font.bold, p.font.color.rgb = Pt(12.5), True, _rgb(ACCENT)
    return tb


def _bullets(shapes, x, y, w, h, lines, size=13.0, sub=11.5):
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (indent, text) in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = ("– " if indent else "• ") + text
        para.level = 1 if indent else 0
        para.font.size = Pt(sub if indent else size)
        para.font.color.rgb = _rgb(MUTED if indent else INK)
        para.space_after = Pt(4)
    return tb


def build() -> str:
    if not CHART_PNG.exists():
        build_chart()

    corr = pd.read_csv(C.ARTIFACTS / "wpb_gwr_correlations.csv")
    prov = json.loads((C.ARTIFACTS / "wpb_gwr_provenance.json").read_text())
    cmp_ = pd.read_csv(C.ARTIFACTS / "wpb_gwr_comparison.csv")
    sub = cmp_[cmp_.n_meters >= 30]
    r100 = corr[(corr.min_meters == 100)
                & (corr.reference == "gwr_separate_wpb_rate")].iloc[0]
    r30 = corr[(corr.min_meters == 30)
               & (corr.reference == "gwr_separate_wpb_rate")].iloc[0]

    flag_rate = 100 * sub.n_flagged.sum() / sub.n_meters.sum()
    sep_rate = 100 * sub.n_separate_wpb.sum() / sub.n_buildings.sum()
    pct_combined = 100 * prov["share_of_7610_that_is_combined"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.26), Inches(12.4), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    p.text = "Wärmepumpenboiler: method, and what the public register can confirm"
    p.font.size, p.font.bold, p.font.color.rgb = Pt(25), True, _rgb(INK)

    tb2 = s.shapes.add_textbox(Inches(0.5), Inches(0.78), Inches(12.4), Inches(0.35))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = ("Label-free physical detection on 15-min grid import, canton Aargau · "
               "January/July 2026 pair · compared against BFS GWR")
    p2.font.size, p2.font.color.rgb = Pt(12.5), _rgb(MUTED)

    # ---- Left column: method -----------------------------------------
    _header(s.shapes, "METHOD", 0.5, 1.28, 4.5)
    _bullets(s.shapes, 0.5, 1.62, 4.5, 5.5, [
        (False, "Per meter-month: 10th percentile (p10) of consumption "
                "across all valid days, per 15-min slot."),
        (True, "a load must repeat on ≥90% of days to lift p10 → timer "
               "regularity is enforced structurally"),
        (False, "Find plateaus ≥0.25 kW above the meter's quietest rolling "
                "hour, wrapping midnight."),
        (False, "WPB candidate only if amplitude 0.25–1.20 kW and duration "
                "1.5–9 h."),
        (True, "bands come from device physics (COP≈3 vs. a resistance "
               "boiler), not fit to labels"),
        (False, "Require a matching plateau in BOTH January and July, "
                "similar amplitude and start time."),
        (True, "rejects seasonal confounders: dehumidifiers, floor heating, "
               "part-load space heat pumps"),
        (False, "Score = 0.5 + amplitude-ratio + start-time-similarity; "
                "flagged at ≥0.8."),
        (False, "No supervised classifier: 4 labelled positives, zero "
                "confirmed negatives."),
    ])

    # ---- Right: chart + reading --------------------------------------
    _header(s.shapes, "CLASSIFIED vs. EXPECTED FROM THE PUBLIC REGISTER",
            5.3, 1.28, 7.6)
    if CHART_PNG.exists():
        s.shapes.add_picture(str(CHART_PNG), Inches(5.25), Inches(1.58),
                             width=Inches(7.75))

    _bullets(s.shapes, 5.3, 4.50, 7.65, 2.7, [
        (False, f"GWR has no Wärmepumpenboiler code. GWAERZW1=7610 is plain "
                f"\"Wärmepumpe\"; {pct_combined:.0f}% of those buildings also "
                f"have a heat pump for space heating — one machine, not a WPB."),
        (False, f"Restricting to hot-water heat pump AND space heating that is "
                f"not a heat pump leaves {sep_rate:.2f}% of buildings, against "
                f"{flag_rate:.2f}% of meters flagged — the same order of magnitude."),
        (False, "Not a validation. The register count is a lower bound (a WPB "
                "behind a space-heating heat pump is absorbed into the "
                f"{pct_combined:.0f}%), buildings are not meters, and the rank "
                f"correlation is only ρ={r100.rho:+.2f} (p={r100.pvalue:.3f}) on "
                f"postcodes with ≥100 meters, ρ={r30.rho:+.2f} (p={r30.pvalue:.2f}) at ≥30."),
    ], size=12.0)

    C.ARTIFACTS.mkdir(parents=True, exist_ok=True)
    prs.save(DECK)
    return str(DECK)


if __name__ == "__main__":
    print("wrote", build())
